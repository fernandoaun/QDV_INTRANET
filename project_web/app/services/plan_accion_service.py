"""Plan de acción (registro SGI QDV-RG-PG-02_03): una fila por acción, editable, con historial."""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any

from sqlalchemy import func, select

from app.extensions import db
from app.models import AppParametro, PlanAccion, PlanAccionCambio
from app.services.fichas_proceso_service import fecha_local
from app.utils.datetime_operacion import now_operacion_naive_local

REGISTRO = {
    "codigo": "QDV-RG-PG-02_03",
    "titulo": "PLAN DE ACCIÓN",
    "fecha_vigencia": date(2026, 9, 22),
    "revision": "00",
}
PARAM_RESPONSABLE = "plan_accion_responsable_seguimiento"
RESPONSABLE_DEFAULT = "Responsable SGC"

# Listas desplegables de la planilla.
OPCIONES: dict[str, tuple[str, ...]] = {
    "origen": (
        "Auditoría interna", "Auditoría de cliente", "Auditoría externa / certificación", "Auditoría externa / vigilancia",
        "Ente regulador", "Evaluación", "Incidente", "Inspección", "Observación preventiva", "Simulacro",
        "Orden de servicio", "Minuta de reunión", "Registro de NC, PNC y OM",
    ),
    "tipo_hallazgo": ("No Conformidad", "Oportunidad de Mejora", "Otros"),
    "tipo_accion": ("Acción corrección (inmediata)", "Acción correctiva", "Acción de control"),
    "causa_raiz": ("Aplica", "No aplica"),
    "estado": ("Cerrado", "En curso", "Atrasado"),
    "resultado_eficacia": ("Eficaz", "No eficaz"),
}
CAMPOS_TEXTO: tuple[str, ...] = (
    "origen", "tipo_hallazgo", "tipo_accion", "accion", "causa_raiz", "causa_raiz_descripcion", "requisito",
    "estado", "responsable", "metodo_eficacia", "resultado_eficacia", "comentarios",
)
CAMPOS_FECHA: tuple[str, ...] = ("fecha_deteccion", "fecha_cumplimiento")
CAMPO_LABELS: dict[str, str] = {
    "numero": "N° acción",
    "origen": "Origen de la acción",
    "tipo_hallazgo": "Tipo de hallazgo resultante de la acción",
    "fecha_deteccion": "Fecha de detección de la acción",
    "tipo_accion": "Tipo de acción",
    "accion": "Acción definida",
    "causa_raiz": "Causa raíz",
    "causa_raiz_descripcion": "Descripción de la causa raíz",
    "requisito": "Requisito / documento asociado a la acción",
    "fecha_cumplimiento": "Fecha de cumplimiento",
    "estado": "Estado de la acción",
    "responsable": "Responsable de cumplimiento",
    "metodo_eficacia": "Método de la eficacia",
    "resultado_eficacia": "Resultado de la eficacia",
    "comentarios": "Comentarios",
    "responsable_seguimiento": "Responsable de seguimiento del plan",
    "alta": "Alta",
    "baja": "Baja",
}
_LARGOS = {"origen": 80, "tipo_hallazgo": 80, "tipo_accion": 80, "causa_raiz": 20, "requisito": 255, "estado": 20,
           "responsable": 255, "resultado_eficacia": 20}


def _today() -> date:
    return now_operacion_naive_local().date()


def fecha_dmy(d: date | None) -> str:
    return d.strftime("%d/%m/%Y") if d else ""


def labels_context() -> dict[str, Any]:
    return {"reg": REGISTRO, "opciones": OPCIONES, "plan_campo_labels": CAMPO_LABELS, "fecha_local": fecha_local,
            "fecha_dmy": fecha_dmy}


def listar() -> list[PlanAccion]:
    return list(db.session.scalars(select(PlanAccion).where(PlanAccion.activo.is_(True)).order_by(PlanAccion.numero)))


def get(accion_id: int) -> PlanAccion | None:
    a = db.session.get(PlanAccion, int(accion_id))
    return a if a is not None and a.activo else None


def responsable_seguimiento() -> str:
    row = db.session.get(AppParametro, PARAM_RESPONSABLE)
    return (row.valor if row is not None else "") or RESPONSABLE_DEFAULT


def guardar_responsable(valor: str, user_id: int | None) -> None:
    nuevo = (valor or "").strip()[:160] or RESPONSABLE_DEFAULT
    row = db.session.get(AppParametro, PARAM_RESPONSABLE) or AppParametro(clave=PARAM_RESPONSABLE)
    row.valor = nuevo
    row.updated_by_user_id = user_id
    db.session.add(row)


def ultima_actualizacion() -> date | None:
    """Fecha de la última carga o cambio de cualquier acción (encabezado del registro)."""
    ult = db.session.scalar(select(func.max(PlanAccionCambio.created_at)))
    if ult is None:
        return None
    from app.utils.datetime_operacion import operacion_zoneinfo

    return ult.astimezone(operacion_zoneinfo()).date()


def cambios(accion: PlanAccion, limit: int = 100) -> list[PlanAccionCambio]:
    return list(
        db.session.scalars(
            select(PlanAccionCambio)
            .where(PlanAccionCambio.accion_id == accion.id)
            .order_by(PlanAccionCambio.created_at.desc(), PlanAccionCambio.id.desc())
            .limit(limit)
        ).unique()
    )


def _log(accion: PlanAccion, user_id: int | None, campo: str, antes: Any, despues: Any) -> None:
    db.session.add(
        PlanAccionCambio(
            accion_id=accion.id,
            user_id=user_id,
            campo=campo,
            antes=None if antes in (None, "") else str(antes),
            despues=None if despues in (None, "") else str(despues),
        )
    )


def parse_fecha(raw: Any) -> date | None:
    """dd/mm/aaaa, aaaa-mm-dd o fecha de Excel. Vacío → None. Inválida → ValueError."""
    if raw in (None, ""):
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    s = str(raw).strip()
    if not s:
        return None
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return date.fromisoformat(s[:10])


def _opcion(campo: str, valor: str) -> str:
    """Normaliza contra la lista desplegable (sin tildes ni mayúsculas ni espacios de más); si no está, queda como vino."""
    def clave(x: str) -> str:
        import unicodedata

        x = unicodedata.normalize("NFKD", x).encode("ascii", "ignore").decode().lower()
        return re.sub(r"[\s/]+", " ", x).strip()

    for op in OPCIONES.get(campo, ()):
        if clave(op) == clave(valor):
            return op
    return valor


def _limpio(data: Any) -> tuple[dict[str, Any], list[str]]:
    v: dict[str, Any] = {}
    errores: list[str] = []
    for k in CAMPOS_TEXTO:
        val = re.sub(r"[ \t]+", " ", str(data.get(k) or "")).strip()
        if k in OPCIONES and val:
            val = _opcion(k, val)
        v[k] = val[: _LARGOS[k]] if k in _LARGOS else val
    for k in CAMPOS_FECHA:
        try:
            v[k] = parse_fecha(data.get(k))
        except ValueError:
            errores.append(f"{CAMPO_LABELS[k]}: fecha inválida, usá dd/mm/aaaa.")
            v[k] = None
    if not v["accion"]:
        errores.append("La acción definida es obligatoria.")
    return v, errores


def guardar(accion: PlanAccion, data: Any, user_id: int | None) -> list[str]:
    v, errores = _limpio(data)
    if errores:
        return errores
    for campo, nuevo in v.items():
        antes = getattr(accion, campo)
        if (antes or None) == (nuevo or None):
            continue
        setattr(accion, campo, nuevo)
        fmt = fecha_dmy if campo in CAMPOS_FECHA else (lambda x: x)
        _log(accion, user_id, campo, fmt(antes) if antes else antes, fmt(nuevo) if nuevo else nuevo)
    return []


def crear(data: Any, user_id: int | None, numero: int | None = None) -> tuple[PlanAccion | None, list[str]]:
    v, errores = _limpio(data)
    if errores:
        return None, errores
    if numero is None:
        numero = (db.session.scalar(select(func.max(PlanAccion.numero))) or 0) + 1
    accion = PlanAccion(numero=numero, **v)
    db.session.add(accion)
    db.session.flush()
    _log(accion, user_id, "alta", None, f"N° {numero}: {accion.accion}")
    return accion, []


def baja(accion: PlanAccion, user_id: int | None) -> None:
    accion.activo = False
    _log(accion, user_id, "baja", f"N° {accion.numero}: {accion.accion}", None)


# ---------------------------------------------------------------- acciones que vienen de otros registros


def sincronizar_filas(filas: list[dict[str, Any]], base: dict[str, Any], user_id: int | None, texto: str = "accion") -> int:
    """Cada renglón con texto y responsable es una acción del plan (minutas, reportes de NC…).

    La primera vez la crea con `base` (origen, requisito, tipo…) y guarda su id en el renglón (`plan_accion_id`);
    después actualiza texto, responsable y fecha de cumplimiento (con historial). Si la acción se dio de baja en el
    plan no se recrea. Modifica `filas` en el lugar y devuelve cuántas acciones creó.
    """
    creadas = 0
    for f in filas:
        if not f.get(texto) or not f.get("responsable"):
            continue
        pid = f.get("plan_accion_id")
        if not pid:
            accion, _ = crear(
                {**base, "accion": f[texto], "estado": "En curso", "responsable": f["responsable"],
                 "fecha_cumplimiento": f.get("fecha")},
                user_id,
            )
            if accion is not None:
                f["plan_accion_id"] = accion.id
                creadas += 1
        elif (accion := get(pid)) is not None:
            datos = {k: getattr(accion, k) for k in (*CAMPOS_TEXTO, *CAMPOS_FECHA)}
            datos.update(accion=f[texto], responsable=f["responsable"], fecha_cumplimiento=f.get("fecha"))
            guardar(accion, datos, user_id)
    return creadas


def conservar_ids(anteriores: list[dict[str, Any]], nuevas: list[dict[str, Any]], texto: str = "accion") -> None:
    """Si un renglón llega sin el id del plan (p. ej. otra pestaña), lo recupera del renglón anterior con el mismo texto."""
    libres = {f.get(texto): f["plan_accion_id"] for f in anteriores if f.get("plan_accion_id")}
    usados = {f["plan_accion_id"] for f in nuevas if f.get("plan_accion_id")}
    for f in nuevas:
        pid = libres.get(f.get(texto))
        if not f.get("plan_accion_id") and pid and pid not in usados:
            f["plan_accion_id"] = pid
            usados.add(pid)


def numeros(ids: list[int]) -> dict[int, int]:
    """id de acción → N° del plan, solo de las acciones activas."""
    return {a.id: a.numero for a in (get(i) for i in ids) if a is not None}


# ---------------------------------------------------------------- lectura de la planilla Excel (.xlsx)

# Columna de la planilla → campo (fila de títulos: la que dice «N° ACCIÓN»).
_COLUMNAS: dict[str, str] = {
    "origen de la accion": "origen",
    "tipo de hallazgo resultante de la accion": "tipo_hallazgo",
    "fecha de deteccion de la accion": "fecha_deteccion",
    "tipo de accion": "tipo_accion",
    "accion definida": "accion",
    "cauza raiz": "causa_raiz",
    "causa raiz": "causa_raiz",
    "descripcion de la causa raiz": "causa_raiz_descripcion",
    "requisito/documento asociado a la accion": "requisito",
    "fecha de cumplimiento": "fecha_cumplimiento",
    "estado de la accion": "estado",
    "responsable de cumplimiento": "responsable",
    "metodo de la eficacia": "metodo_eficacia",
    "resultado de la eficacia": "resultado_eficacia",
    "comentarios": "comentarios",
}


def _norm(x: Any) -> str:
    import unicodedata

    s = unicodedata.normalize("NFKD", str(x or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s).strip()


def leer_xlsx(contenido: bytes) -> tuple[dict[str, Any], list[str]]:
    """Lee la planilla del registro: responsable de seguimiento y una fila por acción con texto.

    Devuelve ({"responsable_seguimiento", "acciones": [{numero, campos…}]}, avisos). Las fechas que no existen
    (p. ej. 31/09) se corrigen al último día del mes y se avisan.
    """
    import calendar

    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True)
    ws = wb.worksheets[0]
    avisos: list[str] = []
    fila_titulos = None
    responsable = ""
    for row in ws.iter_rows(min_row=1, max_row=40):
        for c in row:
            t = _norm(c.value)
            if t.startswith("responsable de seguimiento"):
                vecinas = [ws.cell(c.row, col).value for col in range(c.column + 1, c.column + 4)]
                responsable = next((str(x).strip() for x in vecinas if x not in (None, "")), "")
            if t == "n° accion" or t == "n accion" or t.startswith("n° acc"):
                fila_titulos = c.row
                col_num = c.column
        if fila_titulos:
            break
    if not fila_titulos:
        return {"responsable_seguimiento": responsable, "acciones": []}, ["No encontré la fila de títulos «N° ACCIÓN»."]
    columnas: dict[int, str] = {}
    for c in ws[fila_titulos]:
        campo = _COLUMNAS.get(_norm(c.value))
        if campo:
            columnas[c.column] = campo
    acciones = []
    for r in range(fila_titulos + 1, ws.max_row + 1):
        datos: dict[str, Any] = {}
        for col, campo in columnas.items():
            val = ws.cell(r, col).value
            if campo in CAMPOS_FECHA and isinstance(val, str):
                m = re.fullmatch(r"\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*", val)
                if m:
                    d, mes, a = int(m.group(1)), int(m.group(2)), int(m.group(3))
                    ultimo = calendar.monthrange(a, mes)[1]
                    if d > ultimo:
                        avisos.append(f"Fila {r}: la fecha {val.strip()} no existe; la tomé como {ultimo:02d}/{mes:02d}/{a}.")
                        d = ultimo
                    val = date(a, mes, d)
            datos[campo] = val
        if not str(datos.get("accion") or "").strip():
            continue
        num = ws.cell(r, col_num).value
        try:
            datos["numero"] = int(str(num).strip())
        except (TypeError, ValueError):
            datos["numero"] = None
        acciones.append(datos)
    return {"responsable_seguimiento": responsable, "acciones": acciones}, avisos
