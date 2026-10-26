"""Matriz de riesgos y oportunidades (registro SGI QDV-RG-PG-09_04).

La planilla original tiene una pestaña por sector: acá cada sector es una hoja, editable y con historial.
El valor del riesgo (probabilidad × severidad) y su clasificación se calculan con la tabla de valoración.
"""
from __future__ import annotations

import json
import re
from datetime import date, datetime
from typing import Any

from sqlalchemy import func, select

from app.extensions import db
from app.models import MatrizRiesgos, MatrizRiesgosCambio
from app.services.fichas_proceso_service import fecha_local
from app.services.matriz_legal_service import procesos
from app.services.plan_accion_service import fecha_dmy, parse_fecha

REGISTRO = {
    "codigo": "QDV-RG-PG-09_04",
    "titulo": "MATRIZ DE RIESGOS Y OPORTUNIDADES",
    "fecha_vigencia": date(2026, 7, 21),
    "revision": "00",
}
TIPOS: tuple[str, ...] = ("Económico", "Estratégico", "Operacional", "Ambiental", "Seguridad", "Salud")
CAMPOS_FILA: tuple[str, ...] = (
    "tipo", "riesgo", "probabilidad", "severidad", "acciones", "fecha", "responsables", "eficacia",
)
FILAS_MIN = 10
CAMPO_LABELS: dict[str, str] = {
    "sector": "Sector",
    "filas": "Riesgos / oportunidades",
    "alta": "Alta",
    "baja": "Baja",
}

# Tabla de valoración de la planilla (pestaña «Tabla de valoración»).
SEVERIDAD: tuple[tuple[int, str, str, str, str], ...] = (
    (5, "Crítico",
     "• Fatalidad de uno o más trabajadores.\n• Exposición masiva a sustancias químicas con consecuencias irreversibles.\n"
     "• Incapacidad permanente múltiple.",
     "• Derrame o liberación de sustancias químicas que genere contaminación significativa del suelo o cuerpos de agua.\n"
     "• Incumplimiento legal con sanciones graves o suspensión de actividades.\n• Daño ambiental de difícil recuperación.",
     "• Paralización de la planta por más de 30 días.\n• Pérdida total o masiva de producción.\n"
     "• Retiro de productos del mercado.\n• Clausura parcial o total por autoridad competente.\n"
     "• Riesgo para la continuidad del negocio."),
    (4, "Importante",
     "• Lesión con tiempo perdido.\n• Hospitalización.\n• Intoxicación química con recuperación prolongada.\n"
     "• Incapacidad permanente parcial.",
     "• Derrame importante contenido dentro del predio que requiere remediación.\n"
     "• Generación de efluentes fuera de especificación.\n• Afectación ambiental recuperable mediante acciones correctivas.",
     "• Parada de producción entre 7 y 30 días.\n• Incumplimiento de requisitos del cliente.\n"
     "• Rechazo de lotes de producción.\n• Daño importante a la imagen de la empresa."),
    (3, "Moderado",
     "• Lesión con días de reposo.\n• Exposición química leve con atención médica.\n• Evacuación parcial del sector.",
     "• Derrame menor contenido dentro de la planta.\n• Emisiones o pérdidas que requieren limpieza y disposición controlada.\n"
     "• Desvíos ambientales sin consecuencias permanentes.",
     "• Interrupción de la producción entre 1 y 7 días.\n• Retraso en entregas.\n• Reproceso de producto.\n"
     "• No conformidades internas o de clientes sin impacto mayor."),
    (2, "Menor",
     "• Lesiones leves que requieren atención médica.\n• Irritación por contacto con productos químicos.\n"
     "• Trabajo restringido temporalmente.",
     "• Derrames pequeños controlados inmediatamente.\n"
     "• Generación ocasional de residuos fuera del procedimiento sin impacto significativo.\n• Emisiones puntuales controladas.",
     "• Interrupción menor a un día.\n• Retrasos administrativos u operativos.\n• Desvíos documentales.\n"
     "• Pérdidas económicas menores recuperables."),
    (1, "Insignificante",
     "• Primeros auxilios únicamente.\n• Sin pérdida de jornada laboral.",
     "• Sin impacto ambiental o impacto despreciable completamente controlado.",
     "• Sin interrupción del proceso.\n• Corrección inmediata sin afectar al cliente ni a la producción."),
)
PROBABILIDAD: tuple[tuple[int, str, str], ...] = (
    (1, "Muy baja", "Los controles existentes son los requeridos (*) para la actividad que genera el riesgo/impacto y están "
                    "implementados y verificados (**). El personal está capacitado."),
    (2, "Baja", "Los controles existentes son los requeridos para la tarea y están implementados aunque no verificados. "
                "El personal está capacitado."),
    (3, "Media", "Los controles existentes son los requeridos para la tarea y están parcialmente implementados. "
                 "El personal está parcialmente capacitado."),
    (4, "Alta", "Los controles existentes no son suficientes para la tarea y/o no están implementados. "
                "El personal está parcialmente capacitado."),
    (5, "Muy alta", "No existen controles para la tarea. El personal no conoce la tarea o no está capacitado."),
)
NOTAS_PROBABILIDAD: tuple[str, ...] = (
    "(*) Por requeridos se entiende: requisitos legales, procedimientos, instructivos de la empresa, buenas prácticas, EPP, "
    "capacitación, inspecciones, plan de respuesta ante emergencias, protecciones pasivas, etc.",
    "(**) Por verificado se entiende que se cuenta con evidencia documentada.",
)
# (clasificación, valor desde, valor hasta, tratamiento)
CLASIFICACIONES: tuple[tuple[str, int, int, str], ...] = (
    ("Aceptable", 1, 1,
     "Mantener los controles existentes. El riesgo/impacto es aceptado por la organización, teniendo en cuenta su política y "
     "sus obligaciones. No se toma acción adicional. No es necesaria supervisión mayor a la habitual. Los riesgos no requieren "
     "mayor consideración. Se deben administrar mediante procedimientos de rutina."),
    ("Bajo", 2, 4,
     "Mantener las medidas implementadas y analizar viabilidad (técnica / económica) de medidas adicionales de control de bajo "
     "costo. Se requiere registro y verificación específica para asegurarse que las medidas se aplican en tiempo y forma. Se "
     "deben hacer comprobaciones periódicas para asegurar que los controles se mantienen."),
    ("Moderado", 5, 9,
     "Deben evaluarse, registrarse e implantarse todas aquellas medidas de reducción de riesgo que sean razonablemente "
     "practicables (técnica / económicamente) y que complementen las existentes. Se deben implementar todas las medidas "
     "posibles antes de comenzar o reanudar la tarea. El resto de las medidas debe implementarse en un plazo no mayor a 6 meses. "
     "Se requiere registro y verificación específica para asegurarse que las medidas se aplican en tiempo y forma. Se deben "
     "hacer comprobaciones periódicas para asegurar que los controles se mantienen. Acciones para reducir/aumentar los "
     "riesgos/oportunidades a un nivel tolerable más bajo/alto. Prioridad media."),
    ("Alto", 10, 16,
     "Requiere aplicación de controles adicionales para bajar el nivel de riesgo a moderado o menor como requisito para comenzar "
     "la tarea. Deben evaluarse, registrarse e implantarse las medidas de reducción de riesgo necesarias y que sean "
     "razonablemente posibles hasta reducir el nivel de riesgo a moderado o menor antes de comenzar o reanudar la tarea. Es "
     "posible la aplicación de medidas transitorias en forma inmediata, hasta que se implemente la acción definitiva a corto "
     "plazo (plazo no mayor a 1 mes). Se requiere registro y verificación específica para asegurarse que las medidas se aplican "
     "en tiempo y forma. Se deben hacer comprobaciones periódicas para asegurar que los controles se mantienen. Requieren "
     "medidas de acción con alta prioridad."),
    ("Intolerable", 20, 25,
     "El trabajo no debe comenzar. Suspensión de la actividad. No se podrán aplicar medidas transitorias para continuar la "
     "tarea. Requiere el análisis y la aplicación previa de todas las medidas de control posibles permanentes que disminuyan el "
     "riesgo (a moderado o menor) y que sean razonablemente practicables (técnica / económicamente). Se requiere registro y "
     "verificación específica para asegurarse que las medidas se aplican en tiempo y forma. Se deben hacer comprobaciones "
     "periódicas para asegurar que los controles se mantienen. Requieren acción inmediata y prioridad absoluta en el tratamiento."),
)


def valor(fila: dict[str, Any]) -> int:
    try:
        return int(fila.get("probabilidad") or 0) * int(fila.get("severidad") or 0)
    except (TypeError, ValueError):
        return 0


def clasificacion(v: int) -> str:
    """Igual que la fórmula de la planilla: 1 aceptable, 2-4 bajo, 5-9 moderado, 10-16 alto, más de 16 intolerable."""
    if v <= 0:
        return "Sin evaluar"
    for nombre, _, hasta, _ in CLASIFICACIONES:
        if v <= hasta:
            return nombre
    return "Intolerable"


def filas(m: MatrizRiesgos | None) -> list[dict[str, str]]:
    if m is None:
        return []
    try:
        data = json.loads(m.filas or "[]")
    except ValueError:
        return []
    return [x for x in data if isinstance(x, dict)]


def resumen(m: MatrizRiesgos) -> dict[str, int]:
    """Cantidad de renglones por clasificación (para el listado)."""
    out: dict[str, int] = {}
    for f in filas(m):
        c = clasificacion(valor(f))
        out[c] = out.get(c, 0) + 1
    return out


def labels_context() -> dict[str, Any]:
    return {
        "reg": REGISTRO, "tipos": TIPOS, "mrg_campo_labels": CAMPO_LABELS, "fecha_local": fecha_local, "fecha_dmy": fecha_dmy,
        "filas": filas, "filas_min": FILAS_MIN, "valor": valor, "clasificacion": clasificacion, "resumen": resumen,
        "severidad_tabla": SEVERIDAD, "probabilidad_tabla": PROBABILIDAD, "notas_probabilidad": NOTAS_PROBABILIDAD,
        "clasificaciones": CLASIFICACIONES,
    }


def listar() -> list[MatrizRiesgos]:
    return list(
        db.session.scalars(
            select(MatrizRiesgos).where(MatrizRiesgos.activo.is_(True)).order_by(MatrizRiesgos.orden, MatrizRiesgos.sector)
        )
    )


def get(matriz_id: int) -> MatrizRiesgos | None:
    m = db.session.get(MatrizRiesgos, int(matriz_id))
    return m if m is not None and m.activo else None


def sectores() -> list[str]:
    """Sectores sugeridos: los procesos de las fichas (QDV-RG-PG-09_03)."""
    return procesos()


def cambios(m: MatrizRiesgos, limit: int = 100) -> list[MatrizRiesgosCambio]:
    return list(
        db.session.scalars(
            select(MatrizRiesgosCambio)
            .where(MatrizRiesgosCambio.matriz_id == m.id)
            .order_by(MatrizRiesgosCambio.created_at.desc(), MatrizRiesgosCambio.id.desc())
            .limit(limit)
        ).unique()
    )


def _log(m: MatrizRiesgos, user_id: int | None, campo: str, antes: Any, despues: Any) -> None:
    db.session.add(
        MatrizRiesgosCambio(
            matriz_id=m.id,
            user_id=user_id,
            campo=campo,
            antes=None if antes in (None, "") else str(antes),
            despues=None if despues in (None, "") else str(despues),
        )
    )


def _fecha(txt: str) -> str:
    """Fecha de implementación: la planilla usa mes/año (12-2026) o fecha completa. Devuelve mm/aaaa o dd/mm/aaaa."""
    t = txt.strip()
    if not t:
        return ""
    if mm := re.fullmatch(r"(\d{1,2})\s*[-/.]\s*(\d{4})", t):
        mes, anio = int(mm.group(1)), int(mm.group(2))
        if 1 <= mes <= 12:
            return f"{mes:02d}/{anio}"
        raise ValueError(t)
    return fecha_dmy(parse_fecha(t))


def _filas_form(data: Any, errores: list[str]) -> str:
    """Renglones desde el formulario (filas_<campo>_<i>) o desde una lista; descarta los vacíos."""
    if hasattr(data, "getlist"):
        crudas = [{k: data.get(f"filas_{k}_{i}") for k in CAMPOS_FILA} for i in range(int(data.get("filas_filas") or 0))]
    else:
        crudas = list(data.get("filas") or [])
    out: list[dict[str, str]] = []
    for i, cruda in enumerate(crudas):
        fila = {k: str(cruda.get(k) if cruda.get(k) is not None else "").strip() for k in CAMPOS_FILA}
        for k in ("tipo", "fecha", "responsables"):
            fila[k] = re.sub(r"\s+", " ", fila[k])[:255]
        if not any(fila.values()):
            continue
        fila["tipo"] = next((t for t in TIPOS if t.lower() == fila["tipo"].lower()), fila["tipo"])
        for k in ("probabilidad", "severidad"):
            if fila[k] and fila[k] not in ("1", "2", "3", "4", "5"):
                errores.append(f"Renglón {i + 1}: la {k} va de 1 a 5.")
        try:
            fila["fecha"] = _fecha(fila["fecha"])
        except ValueError:
            errores.append(f"Renglón {i + 1}: fecha de implementación inválida, usá mm/aaaa o dd/mm/aaaa.")
        if not fila["riesgo"]:
            errores.append(f"Renglón {i + 1}: falta el riesgo / oportunidad.")
        out.append(fila)
    return json.dumps(out, ensure_ascii=False)


def _limpio(data: Any) -> tuple[dict[str, Any], list[str]]:
    errores: list[str] = []
    v: dict[str, Any] = {"sector": re.sub(r"\s+", " ", str(data.get("sector") or "")).strip()[:160]}
    v["filas"] = _filas_form(data, errores)
    if not v["sector"]:
        errores.append("El sector es obligatorio.")
    return v, errores


def _sector_repetido(sector: str, excepto: int | None) -> bool:
    q = select(MatrizRiesgos.id).where(MatrizRiesgos.activo.is_(True), func.lower(MatrizRiesgos.sector) == sector.lower())
    if excepto is not None:
        q = q.where(MatrizRiesgos.id != excepto)
    return db.session.scalar(q) is not None


def _fmt(campo: str, val: Any) -> Any:
    if campo != "filas" or not val:
        return val
    lineas = []
    for i, f in enumerate(json.loads(val), 1):
        v = valor(f)
        lineas.append(
            f"{i}. {f.get('riesgo') or '—'} · {f.get('tipo') or '—'} · P{f.get('probabilidad') or '—'}×S{f.get('severidad') or '—'}"
            f" = {v or '—'} {clasificacion(v)} · {f.get('acciones') or '—'} · {f.get('fecha') or '—'} · {f.get('responsables') or '—'}"
            + (f" · Eficacia: {f['eficacia']}" if f.get("eficacia") else "")
        )
    return "\n".join(lineas) or None


def guardar(m: MatrizRiesgos, data: Any, user_id: int | None) -> list[str]:
    v, errores = _limpio(data)
    if not errores and _sector_repetido(v["sector"], m.id):
        errores.append(f"Ya hay una matriz para el sector «{v['sector']}».")
    if errores:
        return errores
    cambio = False
    for campo, nuevo in v.items():
        antes = getattr(m, campo)
        if campo == "filas" and json.loads(antes or "[]") == json.loads(nuevo):
            continue
        if (antes or None) == (nuevo or None):
            continue
        setattr(m, campo, nuevo)
        _log(m, user_id, campo, _fmt(campo, antes), _fmt(campo, nuevo))
        cambio = True
    if cambio:
        m.fecha_actualizacion = datetime.now().date()
    return []


def crear(data: Any, user_id: int | None) -> tuple[MatrizRiesgos | None, list[str]]:
    v, errores = _limpio(data)
    if not errores and _sector_repetido(v["sector"], None):
        errores.append(f"Ya hay una matriz para el sector «{v['sector']}».")
    if errores:
        return None, errores
    orden = (db.session.scalar(select(func.max(MatrizRiesgos.orden))) or 0) + 1
    m = MatrizRiesgos(orden=orden, fecha_actualizacion=data.get("fecha_actualizacion") or datetime.now().date(), **v)
    db.session.add(m)
    db.session.flush()
    _log(m, user_id, "alta", None, f"{m.sector}: {len(filas(m))} renglón(es)")
    return m, []


def baja(m: MatrizRiesgos, user_id: int | None) -> None:
    m.activo = False
    _log(m, user_id, "baja", m.sector, None)


# ---------------------------------------------------------------- importación de la planilla original


def leer_xlsx(fuente: Any) -> list[dict[str, Any]]:
    """Lee la planilla QDV-RG-PG-09_04: una hoja por sector (las pestañas con el código en K1 y el sector en C6).

    Devuelve [{sector, fecha_actualizacion, filas}] listo para `crear`. La pestaña «Tabla de valoración» se saltea.
    """
    import openpyxl

    wb = openpyxl.load_workbook(fuente, data_only=True)
    out: list[dict[str, Any]] = []
    for ws in wb.worksheets:
        if REGISTRO["codigo"] not in str(ws["K1"].value or "") or not str(ws["C6"].value or "").strip():
            continue
        fecha_act = ws["C7"].value
        filas_: list[dict[str, str]] = []
        for row in ws.iter_rows(min_row=12, max_row=ws.max_row, max_col=11, values_only=True):
            _, tipo, riesgo, prob, sev, _, _, acciones, fecha, resp, eficacia = (list(row) + [None] * 11)[:11]
            if not str(riesgo or "").strip():
                continue
            if isinstance(fecha, datetime):
                fecha = f"{fecha.month:02d}/{fecha.year}"
            try:
                fecha = _fecha(str(fecha or ""))
            except ValueError:
                pass
            filas_.append({
                "tipo": str(tipo or "").strip(), "riesgo": str(riesgo).strip(),
                "probabilidad": str(prob or "").strip(), "severidad": str(sev or "").strip(),
                "acciones": str(acciones or "").strip(), "fecha": str(fecha or "").strip(),
                "responsables": str(resp or "").strip(), "eficacia": str(eficacia or "").strip(),
            })
        out.append({
            "sector": str(ws["C6"].value).strip(),
            "fecha_actualizacion": fecha_act.date() if isinstance(fecha_act, datetime) else None,
            "filas": filas_,
        })
    return out
