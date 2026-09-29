"""Programas anuales del PG-02 (registros SGI QDV-RG-PG-02_xx): filas con seguimiento mensual por colores.

Cada `tipo` (Objetivos 02_01, CMASS 02_02, …) comparte el mismo motor y define en `TIPOS` sus columnas,
leyenda de estados y qué campos edita SGI y cuáles solo el administrador.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from typing import Any, BinaryIO

from sqlalchemy import select

from app.extensions import db
from app.models import Objetivo, ObjetivoCambio, ObjetivoMes, ObjetivoPrograma
from app.utils.datetime_operacion import now_operacion_naive_local

COLOR_ENCABEZADO = "FABF8F"
MESES: tuple[str, ...] = ("ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", "NOV", "DIC")
MES_ANUAL = 13

ESTADO_LABELS: dict[str, str] = {
    "realizado": "Realizado",
    "en_implementacion": "En implementación",
    "atrasado": "Atrasado",
    "programado": "Programado",
    "na": "N/A",
}
ESTADO_COLORES: dict[str, str] = {
    "realizado": "92D050",
    "en_implementacion": "FFFF99",
    "atrasado": "FF0000",
    "programado": "808080",
    "na": "FFFFFF",
}

CAMPO_LABELS: dict[str, str] = {
    "proceso": "Proceso",
    "objetivo": "Objetivo",
    "meta": "Meta",
    "indicador": "Indicador",
    "recursos": "Recursos / actividades",
    "frecuencia": "Frecuencia",
    "estado_cumplimiento": "Estado de cumplimiento",
    "observaciones": "Observaciones",
    "alta": "Alta",
    "baja": "Baja",
    "encabezado": "Encabezado",
    "importacion": "Importación de planilla",
}


@dataclass(frozen=True)
class TipoPrograma:
    clave: str
    codigo: str
    titulo: str  # como figura en el encabezado de la planilla
    nombre: str  # como se lo nombra en el sistema
    descripcion: str
    fila: str  # cómo se llama cada fila: «objetivo», «actividad»
    estados: tuple[str, ...]
    campos_admin: tuple[str, ...]
    campos_seguimiento: tuple[str, ...]
    con_anual: bool
    agrupar_proceso: bool
    # Letra que se escribe en la celda del mes al elegir un estado (CMASS usa R/A/P).
    letras: dict[str, str] = field(default_factory=dict)
    frecuencias: tuple[str, ...] = ("Mensual", "Bimestral", "Trimestral", "Semestral", "Anual")


TIPOS: dict[str, TipoPrograma] = {
    "objetivos": TipoPrograma(
        clave="objetivos",
        codigo="QDV-RG-PG-02_01",
        titulo="PROGRAMA DE OBJETIVOS",
        nombre="Programa de Objetivos",
        descripcion="Objetivos, metas e indicadores con seguimiento mensual.",
        fila="objetivo",
        estados=("realizado", "en_implementacion", "atrasado", "na"),
        campos_admin=("proceso", "objetivo", "meta", "indicador"),
        campos_seguimiento=("recursos", "frecuencia", "estado_cumplimiento"),
        con_anual=True,
        agrupar_proceso=False,
    ),
    "cmass": TipoPrograma(
        clave="cmass",
        codigo="QDV-RG-PG-02_02",
        titulo="PROGRAMA CMASS",
        nombre="Programa CMASS",
        descripcion="Actividades de calidad, medio ambiente, seguridad y salud con seguimiento mensual.",
        fila="actividad",
        estados=("realizado", "atrasado", "programado"),
        campos_admin=("proceso",),
        campos_seguimiento=("objetivo", "frecuencia", "observaciones"),
        con_anual=False,
        agrupar_proceso=True,
        letras={"realizado": "R", "atrasado": "A", "programado": "P"},
        frecuencias=("Mensual", "Bimestral", "Trimestral", "Semestral", "Anual", "Unica vez"),
    ),
}


def tipo_o_none(clave: str | None) -> TipoPrograma | None:
    return TIPOS.get((clave or "").strip())


def _today() -> date:
    return now_operacion_naive_local().date()


def fecha_local(dt: datetime | None) -> str:
    """Fecha y hora de un cambio (guardado en UTC) en hora de planta."""
    if dt is None:
        return ""
    from app.utils.datetime_operacion import operacion_zoneinfo

    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(operacion_zoneinfo()).strftime("%d/%m/%Y %H:%M")


def mes_label(mes: int) -> str:
    return "ANUAL" if mes == MES_ANUAL else MESES[mes - 1]


def ultimo_mes(t: TipoPrograma) -> int:
    return MES_ANUAL if t.con_anual else 12


def labels_context(t: TipoPrograma) -> dict[str, Any]:
    return {
        "tp": t,
        "obj_estados": t.estados,
        "obj_estado_labels": ESTADO_LABELS,
        "obj_estado_colores": ESTADO_COLORES,
        "obj_letras": t.letras,
        "obj_meses": MESES,
        "obj_mes_anual": MES_ANUAL,
        "obj_ultimo_mes": ultimo_mes(t),
        "obj_frecuencias": t.frecuencias,
        "obj_color_encabezado": COLOR_ENCABEZADO,
        "obj_campo_labels": {**CAMPO_LABELS, "objetivo": t.fila.capitalize()},
        "obj_mes_label": mes_label,
        "fecha_local": fecha_local,
    }


def grupos_por_proceso(objs: list[Objetivo]) -> dict[int, int]:
    """id de la primera fila de cada tramo consecutivo con el mismo proceso → cantidad de filas (rowspan)."""
    out: dict[int, int] = {}
    inicio: Objetivo | None = None
    for o in objs:
        if inicio is not None and o.proceso == inicio.proceso:
            out[inicio.id] += 1
        else:
            inicio = o
            out[o.id] = 1
    return out


# ---------------------------------------------------------------- consultas


def get_programa(tipo: str, anio: int) -> ObjetivoPrograma | None:
    return db.session.scalar(
        select(ObjetivoPrograma).where(ObjetivoPrograma.tipo == tipo, ObjetivoPrograma.anio == int(anio))
    )


def anios_disponibles(tipo: str) -> list[int]:
    return list(
        db.session.scalars(
            select(ObjetivoPrograma.anio).where(ObjetivoPrograma.tipo == tipo).order_by(ObjetivoPrograma.anio.desc())
        )
    )


def anio_por_defecto(tipo: str) -> int:
    """Año en curso si existe su programa; si no, el más reciente cargado."""
    hoy = _today().year
    anios = anios_disponibles(tipo)
    if not anios or hoy in anios:
        return hoy
    return anios[0]


def objetivos_activos(prog: ObjetivoPrograma) -> list[Objetivo]:
    return [o for o in prog.objetivos if o.activo]


def get_objetivo(objetivo_id: int) -> Objetivo | None:
    return db.session.get(Objetivo, int(objetivo_id))


def meses_map(obj: Objetivo) -> dict[int, ObjetivoMes]:
    return {int(m.mes): m for m in obj.meses}


def cambios_recientes(prog: ObjetivoPrograma, limit: int = 200) -> list[ObjetivoCambio]:
    stmt = (
        select(ObjetivoCambio)
        .where(ObjetivoCambio.programa_id == prog.id)
        .order_by(ObjetivoCambio.created_at.desc(), ObjetivoCambio.id.desc())
        .limit(limit)
    )
    return list(db.session.scalars(stmt).unique())


# ---------------------------------------------------------------- escritura


def _log(prog: ObjetivoPrograma, obj: Objetivo | None, user_id: int | None, campo: str, antes: Any, despues: Any) -> None:
    db.session.add(
        ObjetivoCambio(
            programa_id=prog.id,
            objetivo_id=obj.id if obj is not None else None,
            user_id=user_id,
            campo=campo,
            antes=None if antes is None else str(antes),
            despues=None if despues is None else str(despues),
        )
    )
    prog.fecha_actualizacion = _today()


def crear_programa(
    tipo: str, anio: int, user_id: int | None, copiar_de: ObjetivoPrograma | None = None
) -> ObjetivoPrograma:
    """Crea el programa del año. Con `copiar_de`, copia las filas activas no cumplidas (sin seguimiento)."""
    t = TIPOS[tipo]
    prog = ObjetivoPrograma(tipo=tipo, anio=int(anio), codigo=t.codigo, revision="00", fecha_actualizacion=_today())
    if copiar_de is not None:
        prog.codigo = copiar_de.codigo
        prog.revision = copiar_de.revision
        prog.fecha_vigencia = copiar_de.fecha_vigencia
        orden = 0
        for o in objetivos_activos(copiar_de):
            if o.estado_cumplimiento == "realizado":
                continue
            orden += 1
            prog.objetivos.append(
                Objetivo(
                    orden=orden,
                    proceso=o.proceso,
                    objetivo=o.objetivo,
                    meta=o.meta,
                    indicador=o.indicador,
                    recursos=o.recursos,
                    frecuencia=o.frecuencia,
                )
            )
    db.session.add(prog)
    db.session.flush()
    _log(prog, None, user_id, "alta", None, f"Programa {anio}" + (f" (copiado de {copiar_de.anio})" if copiar_de else ""))
    return prog


def _clean(t: TipoPrograma, campo: str, raw: Any) -> tuple[Any, str | None]:
    v = (raw or "").strip() if isinstance(raw, str) or raw is None else raw
    if campo == "estado_cumplimiento":
        if not v:
            return None, None
        if v not in t.estados:
            return None, "Estado de cumplimiento inválido."
        return v, None
    if campo == "proceso":
        if not v:
            return None, "El proceso es obligatorio."
        return v[:128], None
    if campo == "objetivo":
        if not v:
            return None, f"La descripción del {t.fila} es obligatoria."
        return v, None
    if campo in ("meta", "indicador"):
        return v[:256], None
    if campo == "frecuencia":
        return (v or "Mensual")[:64], None
    return v, None


def _tipo_de(obj: Objetivo) -> TipoPrograma:
    return TIPOS[obj.programa.tipo]


def actualizar_campos(obj: Objetivo, data: dict[str, Any], campos: tuple[str, ...], user_id: int | None) -> list[str]:
    """Aplica solo los `campos` permitidos presentes en `data`. Devuelve errores (sin aplicar nada si hay)."""
    t = _tipo_de(obj)
    nuevos: dict[str, Any] = {}
    errores: list[str] = []
    for campo in campos:
        if campo not in data:
            continue
        v, err = _clean(t, campo, data.get(campo))
        if err:
            errores.append(err)
        else:
            nuevos[campo] = v
    if errores:
        return errores
    prog = obj.programa
    for campo, v in nuevos.items():
        antes = getattr(obj, campo)
        if (antes or None) == (v or None):
            continue
        setattr(obj, campo, v if v is not None or campo == "estado_cumplimiento" else "")
        _log(prog, obj, user_id, campo, antes, v)
    return []


def alta_objetivo(prog: ObjetivoPrograma, data: dict[str, Any], user_id: int | None) -> tuple[Objetivo | None, list[str]]:
    t = TIPOS[prog.tipo]
    obj = Objetivo(orden=max([o.orden for o in prog.objetivos] or [0]) + 1)
    errores: list[str] = []
    for campo in t.campos_admin + t.campos_seguimiento:
        v, err = _clean(t, campo, data.get(campo))
        if err:
            errores.append(err)
        else:
            setattr(obj, campo, v if v is not None or campo == "estado_cumplimiento" else "")
    if errores:
        return None, errores
    # Una actividad nueva se ubica al final de su proceso para no partir el agrupamiento.
    if t.agrupar_proceso:
        mismos = [o.orden for o in prog.objetivos if o.proceso == obj.proceso]
        if mismos:
            obj.orden = max(mismos) + 1
            for o in prog.objetivos:
                if o.orden >= obj.orden:
                    o.orden += 1
    prog.objetivos.append(obj)
    prog.objetivos.sort(key=lambda o: o.orden)
    db.session.flush()
    _log(prog, obj, user_id, "alta", None, obj.objetivo[:200])
    return obj, []


def baja_objetivo(obj: Objetivo, user_id: int | None) -> None:
    if not obj.activo:
        return
    obj.activo = False
    _log(obj.programa, obj, user_id, "baja", obj.objetivo[:200], None)


def _mes_texto(estado: str | None, valor: str) -> str:
    partes = [ESTADO_LABELS.get(estado or "", "")] if estado else []
    if valor:
        partes.append(valor)
    return " · ".join(partes) or "(vacío)"


def set_mes(
    obj: Objetivo, mes: int, estado: str | None, valor: str | None, user_id: int | None
) -> tuple[ObjetivoMes | None, str | None]:
    t = _tipo_de(obj)
    if not 1 <= int(mes) <= ultimo_mes(t):
        return None, "Mes inválido."
    estado = (estado or "").strip() or None
    if estado is not None and estado not in t.estados:
        return None, "Estado inválido."
    valor = (valor or "").strip()[:64]
    m = meses_map(obj).get(int(mes))
    antes = _mes_texto(m.estado, m.valor) if m is not None else "(vacío)"
    if m is None:
        m = ObjetivoMes(mes=int(mes), estado=estado, valor=valor)
        obj.meses.append(m)
    elif m.estado == estado and (m.valor or "") == valor:
        return m, None
    else:
        m.estado = estado
        m.valor = valor
    _log(obj.programa, obj, user_id, f"mes:{mes_label(int(mes))}", antes, _mes_texto(estado, valor))
    return m, None


def actualizar_encabezado(prog: ObjetivoPrograma, data: dict[str, Any], user_id: int | None) -> list[str]:
    codigo = (data.get("codigo") or "").strip()[:64]
    revision = (data.get("revision") or "").strip()[:16]
    raw_vig = (data.get("fecha_vigencia") or "").strip()
    if not codigo or not revision:
        return ["Código y revisión son obligatorios."]
    vig: date | None = None
    if raw_vig:
        try:
            vig = date.fromisoformat(raw_vig)
        except ValueError:
            return ["Fecha de vigencia inválida."]
    antes = f"{prog.codigo} · Rev. {prog.revision} · {prog.fecha_vigencia or '—'}"
    prog.codigo, prog.revision, prog.fecha_vigencia = codigo, revision, vig
    despues = f"{prog.codigo} · Rev. {prog.revision} · {prog.fecha_vigencia or '—'}"
    if antes != despues:
        _log(prog, None, user_id, "encabezado", antes, despues)
    return []


# ---------------------------------------------------------------- importación desde la planilla

# Encabezados de columna de la planilla → campo de la fila.
_COLUMNAS_PLANILLA: dict[str, str] = {
    "OBJETIVO": "objetivo",
    "ACTIVIDADES": "objetivo",
    "ACTIVIDAD": "objetivo",
    "META": "meta",
    "INDICADOR": "indicador",
    "ESTADO DE CUMPLIMIENTO": "estado_cumplimiento",
    "RECURSOS/ACTIVIDADES A DESARROLLAR": "recursos",
    "FRECUENCIA": "frecuencia",
    "OBSERVACIONES": "observaciones",
}


def _color_a_estado(t: TipoPrograma, rgb: str | None) -> str | None:
    if not rgb or len(rgb) < 6:
        return None
    c = rgb[-6:].upper()
    for estado in t.estados:
        if estado != "na" and c == ESTADO_COLORES[estado]:
            return estado
    # Variantes de la paleta estándar de Excel.
    candidatos = {
        "en_implementacion": ("FFFF00", "FFFF66", "FFFFCC", "FFE699", "FFD966"),
        "realizado": ("00B050", "00FF00", "A9D08E", "C6EFCE"),
        "atrasado": ("C00000", "FF5050", "FFC7CE"),
        "programado": ("A6A6A6", "BFBFBF", "7F7F7F", "D9D9D9", "969696"),
    }
    for estado, colores in candidatos.items():
        if estado in t.estados and c in colores:
            return estado
    return None


def _texto_a_estado(t: TipoPrograma, v: Any) -> str | None:
    s = str(v or "").strip().lower()
    for estado in t.estados:
        if s == ESTADO_LABELS[estado].lower():
            return estado
    return None


def _fecha_en_texto(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", str(v or ""))
    if not m:
        return None
    try:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


def _norm(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().upper()


def importar_planilla(fh: BinaryIO, tipo: str, anio: int, user_id: int | None) -> tuple[ObjetivoPrograma | None, list[str]]:
    """Carga un programa desde la planilla Excel del registro. Solo si ese año todavía no tiene filas."""
    import openpyxl

    t = TIPOS[tipo]
    try:
        wb = openpyxl.load_workbook(fh, data_only=True)
    except Exception:
        return None, [f"No se pudo leer el archivo. Tiene que ser la planilla .xlsx del {t.nombre}."]
    ws = wb.active

    fila_enc = next((r for r in range(1, min(ws.max_row, 40) + 1) if _norm(ws.cell(r, 1).value) == "PROCESO"), None)
    if fila_enc is None:
        return None, ["No encontré la fila de encabezados (PROCESO …)."]

    # Fila de meses (ENE…DIC) y columnas de cada campo, buscando en las filas del encabezado.
    fila_meses = None
    col_mes: dict[int, int] = {}
    col_campo: dict[str, int] = {"proceso": 1}
    for r in range(fila_enc, fila_enc + 4):
        for c in range(1, ws.max_column + 1):
            txt = _norm(ws.cell(r, c).value)
            if txt in MESES:
                col_mes[MESES.index(txt) + 1] = c
                fila_meses = r
            elif txt == "ANUAL" and t.con_anual:
                col_mes[MES_ANUAL] = c
            elif txt in _COLUMNAS_PLANILLA and _COLUMNAS_PLANILLA[txt] not in col_campo:
                col_campo[_COLUMNAS_PLANILLA[txt]] = c
    if fila_meses is None or len([m for m in col_mes if m <= 12]) < 12:
        return None, ["No encontré las columnas de meses (ENE a DIC)."]
    if "objetivo" not in col_campo:
        return None, [f"No encontré la columna de {t.fila}es."]

    prog = get_programa(tipo, anio)
    if prog is not None and objetivos_activos(prog):
        return None, [f"El {t.nombre} {anio} ya tiene filas cargadas; no se importa encima."]

    filas: list[dict[str, Any]] = []
    proceso_actual = ""
    for r in range(fila_meses + 1, ws.max_row + 1):
        proceso = str(ws.cell(r, 1).value or "").strip()
        texto = str(ws.cell(r, col_campo["objetivo"]).value or "").strip()
        if not proceso and not texto:
            break
        # Proceso en celdas combinadas: el valor está solo en la primera fila del grupo.
        if proceso:
            proceso_actual = proceso
        fila: dict[str, Any] = {"proceso": (proceso_actual or "—")[:128], "objetivo": texto or "—"}
        for campo, c in col_campo.items():
            if campo in ("proceso", "objetivo"):
                continue
            val = ws.cell(r, c).value
            if campo == "estado_cumplimiento":
                fila[campo] = _texto_a_estado(t, val)
            else:
                fila[campo] = str(val or "").strip()
        fila["frecuencia"] = (fila.get("frecuencia") or "Mensual")[:64]
        fila["meta"] = (fila.get("meta") or "")[:256]
        fila["indicador"] = (fila.get("indicador") or "")[:256]
        meses: dict[int, tuple[str | None, str]] = {}
        for mes, c in col_mes.items():
            cell = ws.cell(r, c)
            fg = cell.fill.fgColor if cell.fill is not None and cell.fill.fill_type == "solid" else None
            estado = _color_a_estado(t, fg.rgb if fg is not None and fg.type == "rgb" else None)
            valor = str(cell.value).strip() if cell.value is not None else ""
            if estado is None and valor and t.letras:
                # Letra sin color (p. ej. «P» sin relleno): se interpreta por la letra.
                estado = next((e for e, l in t.letras.items() if l == valor.upper()), None)
            if estado or valor:
                meses[mes] = (estado, valor[:64])
        fila["meses"] = meses
        filas.append(fila)
    if not filas:
        return None, [f"La planilla no tiene {t.fila}es debajo de los encabezados."]

    if prog is None:
        prog = crear_programa(tipo, anio, user_id)
    for rr in range(1, fila_enc + 1):
        for cc in range(1, ws.max_column + 1):
            txt = str(ws.cell(rr, cc).value or "").strip()
            if not txt:
                continue
            if txt.upper().startswith("QDV-RG"):
                prog.codigo = txt.split()[0][:64]
            elif txt.lower().startswith("fecha de vigencia"):
                prog.fecha_vigencia = _fecha_en_texto(txt) or prog.fecha_vigencia
            elif re.match(r"(?i)^rev\.?\s*\d+", txt):
                prog.revision = re.sub(r"(?i)^rev\.?\s*", "", txt)[:16]
    for i, f in enumerate(filas, start=1):
        obj = Objetivo(
            orden=i,
            proceso=f["proceso"],
            objetivo=f["objetivo"],
            meta=f.get("meta", ""),
            indicador=f.get("indicador", ""),
            estado_cumplimiento=f.get("estado_cumplimiento"),
            recursos=f.get("recursos", ""),
            frecuencia=f["frecuencia"],
            observaciones=f.get("observaciones", ""),
        )
        for mes, (estado, valor) in sorted(f["meses"].items()):
            obj.meses.append(ObjetivoMes(mes=mes, estado=estado, valor=valor))
        prog.objetivos.append(obj)
    db.session.flush()
    _log(prog, None, user_id, "importacion", None, f"{len(filas)} fila(s)")
    return prog, []


# ---------------------------------------------------------------- exportación con el formato del registro


def exportar_xlsx(prog: ObjetivoPrograma) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    t = TIPOS[prog.tipo]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (prog.codigo or t.codigo).replace("QDV-RG-", "QDV-")[:31]
    medium = Side(style="medium", color="000000")
    thin = Side(style="thin", color="000000")
    caja = Border(left=medium, right=medium, top=medium, bottom=medium)
    celda = Border(left=thin, right=thin, top=thin, bottom=thin)
    naranja = PatternFill("solid", fgColor=COLOR_ENCABEZADO)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    izq = Alignment(horizontal="left", vertical="center", wrap_text=True)
    negrita = Font(bold=True, size=12)
    normal = Font(size=12)

    # Columnas de datos antes de los meses, según el tipo de programa.
    if t.clave == "objetivos":
        cols = [
            ("PROCESO", "proceso", 32.5),
            ("OBJETIVO", "objetivo", 37.3),
            ("META", "meta", 26.5),
            ("INDICADOR", "indicador", 36.8),
            ("ESTADO DE CUMPLIMIENTO", "estado_cumplimiento", 27.7),
            ("RECURSOS/ACTIVIDADES A DESARROLLAR", "recursos", 41.2),
            ("FRECUENCIA", "frecuencia", 41.3),
        ]
        cols_post: list[tuple[str, str, float]] = []
    else:
        cols = [("PROCESO", "proceso", 16.5), ("ACTIVIDADES", "objetivo", 70.0), ("FRECUENCIA", "frecuencia", 15.0)]
        cols_post = [("OBSERVACIONES", "observaciones", 40.0)]
    n_meses = ultimo_mes(t)
    c_mes0 = len(cols) + 1
    c_post0 = c_mes0 + n_meses
    ultima = c_post0 + len(cols_post) - 1
    for i, (_, _, w) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for m in range(n_meses):
        ws.column_dimensions[get_column_letter(c_mes0 + m)].width = 8.5
    for i, (_, _, w) in enumerate(cols_post):
        ws.column_dimensions[get_column_letter(c_post0 + i)].width = w
    L = get_column_letter

    # Encabezado del registro: logo | título | código, vigencia, revisión, página.
    ws["A1"] = "QUÍMICA DEL VALLE S.R.L."
    ws["A1"].font = Font(bold=True, size=14, color="E46C0A")
    ws["A1"].alignment = centro
    ws.merge_cells("A1:A4")
    c_meta = max(ultima - 3, 3)
    ws["B1"] = t.titulo
    ws["B1"].font = Font(bold=True, size=20)
    ws["B1"].alignment = centro
    ws.merge_cells(f"B1:{L(c_meta - 1)}4")
    vig = prog.fecha_vigencia.strftime("%d/%m/%Y") if prog.fecha_vigencia else "—"
    for i, txt in enumerate([prog.codigo, f"Fecha de Vigencia: {vig}", f"Rev. {prog.revision}", "Página 1 de 1"], start=1):
        ws.cell(i, c_meta, txt).font = negrita
        ws.cell(i, c_meta).alignment = centro
        ws.merge_cells(start_row=i, start_column=c_meta, end_row=i, end_column=ultima)
    for r in range(1, 5):
        for c in range(1, ultima + 1):
            ws.cell(r, c).border = caja
    act = prog.fecha_actualizacion.strftime("%d/%m/%Y") if prog.fecha_actualizacion else "—"
    ws["A6"] = f"Fecha de actualización: {act}"
    ws["A6"].font = negrita
    ws["A6"].fill = naranja
    ws.merge_cells(f"A6:{L(ultima)}6")

    for i, (titulo, _, _) in enumerate(cols, start=1):
        ws.cell(7, i, titulo)
        ws.merge_cells(start_row=7, start_column=i, end_row=8, end_column=i)
    ws.cell(7, c_mes0, f"AÑO {prog.anio}" if t.clave != "objetivos" else "SEGUIMIENTO / VALOR DEL INDICADOR")
    ws.merge_cells(start_row=7, start_column=c_mes0, end_row=7, end_column=c_post0 - 1)
    for m in range(1, n_meses + 1):
        ws.cell(8, c_mes0 + m - 1, mes_label(m))
    for i, (titulo, _, _) in enumerate(cols_post):
        ws.cell(7, c_post0 + i, titulo)
        ws.merge_cells(start_row=7, start_column=c_post0 + i, end_row=8, end_column=c_post0 + i)
    for r in (7, 8):
        for c in range(1, ultima + 1):
            cell = ws.cell(r, c)
            cell.fill = naranja
            cell.font = negrita
            cell.alignment = centro
            cell.border = caja
    ws.row_dimensions[8].height = 30

    r = 9
    objs = objetivos_activos(prog)
    grupos = grupos_por_proceso(objs) if t.agrupar_proceso else {}
    for o in objs:
        mm = meses_map(o)
        for i, (_, campo, _) in enumerate(cols + cols_post, start=1):
            c = i if i <= len(cols) else c_post0 + (i - len(cols) - 1)
            v = getattr(o, campo) or ""
            if campo == "estado_cumplimiento":
                v = ESTADO_LABELS.get(v, "")
            if campo == "proceso" and t.agrupar_proceso and o.id not in grupos:
                v = None
            cell = ws.cell(r, c, v)
            cell.font = negrita if campo in ("proceso", "frecuencia") else normal
            cell.alignment = izq if campo in ("objetivo", "recursos", "observaciones") else centro
            cell.border = celda
            if campo == "estado_cumplimiento" and o.estado_cumplimiento:
                cell.fill = PatternFill("solid", fgColor=ESTADO_COLORES[o.estado_cumplimiento])
        if t.agrupar_proceso and grupos.get(o.id, 1) > 1:
            ws.merge_cells(start_row=r, start_column=1, end_row=r + grupos[o.id] - 1, end_column=1)
        for mes in range(1, n_meses + 1):
            cell = ws.cell(r, c_mes0 + mes - 1)
            cell.border = celda
            cell.alignment = centro
            cell.font = negrita
            m = mm.get(mes)
            if m is None:
                continue
            if m.valor:
                cell.value = m.valor
            if m.estado:
                cell.fill = PatternFill("solid", fgColor=ESTADO_COLORES[m.estado])
        ws.row_dimensions[r].height = 96 if t.clave == "objetivos" else 20
        r += 1

    r += 2
    for estado in t.estados:
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=ESTADO_COLORES[estado])
        ws.cell(r, 1).border = celda
        letra = t.letras.get(estado)
        ws.cell(r, 2, ESTADO_LABELS[estado] + (f" ({letra})" if letra else "")).alignment = Alignment(vertical="center")
        r += 1

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
