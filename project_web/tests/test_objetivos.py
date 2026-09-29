from __future__ import annotations

import io
import re

import pytest
from werkzeug.security import generate_password_hash


def _planilla() -> bytes:
    """Planilla con la misma estructura que el registro QDV-RG-PG-02_01 (datos ficticios)."""
    import openpyxl
    from openpyxl.styles import PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B1"] = "PROGRAMA DE OBJETIVOS"
    ws["P1"] = "QDV-RG-PG-02_01        "
    ws["P2"] = "Fecha de Vigencia: 26/5/2026"
    ws["P3"] = "Rev. 00"
    ws["A6"] = "Fecha de actualización: 26/5/2026"
    for col, t in zip("ABCDE", ("PROCESO", "OBJETIVO", "META", "INDICADOR", "ESTADO DE CUMPLIMIENTO")):
        ws[f"{col}7"] = t
    ws["F8"], ws["G8"] = "RECURSOS/ACTIVIDADES A DESARROLLAR", "FRECUENCIA"
    for i, m in enumerate(("ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL ", "AGO", "SEP", "OCT", "NOV", "DIC")):
        ws.cell(8, 8 + i, m)
    ws["T8"] = "ANUAL"
    rojo = PatternFill("solid", fgColor="FFFF0000")
    verde = PatternFill("solid", fgColor="FF92D050")
    for r, proc in ((9, "Producción"), (10, "Calidad")):
        ws.cell(r, 1, proc)
        ws.cell(r, 2, f"Objetivo de {proc}")
        ws.cell(r, 3, "100%")
        ws.cell(r, 4, "% avance")
        ws.cell(r, 6, "• Tarea uno\n• Tarea dos")
        ws.cell(r, 7, "Mensual")
        for c in range(8, 16):
            ws.cell(r, c).fill = rojo
    ws.cell(10, 5, "En implementación")
    ws.cell(10, 16).fill = verde
    ws.cell(10, 16, "40 %")
    ws["A18"].fill = verde
    ws["B18"] = "Realizado"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_client(app, client, username, **user_kw):
    from app.extensions import db
    from app.models import User

    with app.app_context():
        db.session.add(User(username=username, password_hash=generate_password_hash("pw-objetivos"), activo=True, **user_kw))
        db.session.commit()
    tok = re.search(r'name="csrf_token"\s+value="([^"]+)"', client.get("/login").get_data(as_text=True)).group(1)
    r = client.post("/login", data={"username": username, "password": "pw-objetivos", "csrf_token": tok})
    assert r.status_code in (302, 303)
    return client


@pytest.fixture
def admin_client(app, client):
    return _make_client(app, client, "pytest_obj_admin", is_admin=True, rol="operaciones")


def _importar(c, anio=2026):
    return c.post(
        f"/sgi/objetivos/{anio}/importar",
        data={"archivo": (io.BytesIO(_planilla()), "programa.xlsx")},
        content_type="multipart/form-data",
    )


def _objetivos(app, anio=2026):
    from app.models import ObjetivoPrograma
    from app.services import objetivos_service as svc

    with app.app_context():
        prog = svc.get_programa(anio)
        assert isinstance(prog, ObjetivoPrograma)
        return prog.codigo, prog.revision, prog.fecha_vigencia, [
            (o.id, o.proceso, o.estado_cumplimiento, o.recursos, {m.mes: (m.estado, m.valor) for m in o.meses})
            for o in svc.objetivos_activos(prog)
        ]


def test_importar_planilla_respeta_colores_y_encabezado(admin_client, app):
    from datetime import date

    r = _importar(admin_client)
    assert r.status_code in (302, 303)
    codigo, rev, vig, objs = _objetivos(app)
    assert (codigo, rev, vig) == ("QDV-RG-PG-02_01", "00", date(2026, 5, 26))
    assert [o[1] for o in objs] == ["Producción", "Calidad"]
    # ENE–AGO en rojo (atrasado), como en la planilla.
    assert all(objs[0][4][m] == ("atrasado", "") for m in range(1, 9))
    assert 9 not in objs[0][4]
    assert objs[1][2] == "en_implementacion"
    assert objs[1][4][9] == ("realizado", "40 %")
    # No se importa dos veces encima.
    _importar(admin_client)
    assert len(_objetivos(app)[3]) == 2

    page = admin_client.get("/sgi/objetivos/2026").get_data(as_text=True)
    assert "PROGRAMA DE OBJETIVOS" in page and "Objetivo de Calidad" in page and "#FF0000" in page


def test_seguimiento_mes_y_campos_admin(admin_client, app):
    _importar(admin_client)
    oid = _objetivos(app)[3][0][0]
    r = admin_client.post(f"/sgi/objetivos/objetivo/{oid}/mes", json={"mes": 9, "estado": "en_implementacion", "valor": "55 %"})
    assert r.status_code == 200 and r.get_json()["ok"] is True
    assert _objetivos(app)[3][0][4][9] == ("en_implementacion", "55 %")
    assert admin_client.post(f"/sgi/objetivos/objetivo/{oid}/mes", json={"mes": 14, "estado": "na"}).status_code == 400

    r = admin_client.post(
        f"/sgi/objetivos/objetivo/{oid}/editar",
        data={"proceso": "Producción / Ambiente", "objetivo": "Nuevo texto", "recursos": "• Otra", "frecuencia": "Mensual"},
    )
    assert r.status_code in (302, 303)
    assert _objetivos(app)[3][0][1] == "Producción / Ambiente"

    hist = admin_client.get("/sgi/objetivos/2026/historial").get_data(as_text=True)
    assert "Seguimiento SEP" in hist and "55 %" in hist

    x = admin_client.get("/sgi/objetivos/2026/export.xlsx")
    assert x.status_code == 200 and x.data[:2] == b"PK"


def test_sgi_edita_tareas_y_colores_pero_no_el_objetivo(app, client):
    from app.extensions import db
    from app.services import objetivos_service as svc

    with app.app_context():
        prog = svc.crear_programa(2026, None)
        obj, errs = svc.alta_objetivo(prog, {"proceso": "Calidad", "objetivo": "Original"}, None)
        db.session.commit()
        oid = obj.id
    c = _make_client(app, client, "pytest_obj_sgi", is_admin=False, rol="sgi")
    r = c.post(f"/sgi/objetivos/objetivo/{oid}/editar", data={"objetivo": "Cambiado", "recursos": "• Tarea SGI"})
    assert r.status_code in (302, 303)
    _, _, _, objs = _objetivos(app)
    assert objs[0][3] == "• Tarea SGI"
    with app.app_context():
        assert svc.get_objetivo(oid).objetivo == "Original"
    assert c.post(f"/sgi/objetivos/objetivo/{oid}/mes", json={"mes": 1, "estado": "realizado"}).get_json()["ok"] is True
    # Alta y baja son del administrador.
    c.post("/sgi/objetivos/2026/objetivos", data={"proceso": "X", "objetivo": "Y"})
    c.post(f"/sgi/objetivos/objetivo/{oid}/baja")
    assert len(_objetivos(app)[3]) == 1


def test_sin_acceso_sgi_no_ve_ni_edita(app, client):
    c = _make_client(app, client, "pytest_obj_oper", is_admin=False, rol="operaciones")
    assert c.get("/sgi/objetivos/2026", follow_redirects=False).status_code in (302, 303)
    assert c.post("/sgi/objetivos/objetivo/1/mes", json={"mes": 1, "estado": "realizado"}).status_code == 403


def test_registro_sgi_puede_asociarse_al_modulo(app):
    from app.services import sgi_procedimiento_service as proc

    with app.app_context():
        links = proc.registro_modulo_links("objetivos")
    assert links["label"] == "Programa de Objetivos"
    assert links["blank_url"] == "/sgi/objetivos/en-blanco"
    assert links["filled_url"] == "/sgi/objetivos/"


def test_ver_en_blanco_no_muestra_datos(admin_client, app):
    _importar(admin_client)
    html = admin_client.get("/sgi/objetivos/en-blanco").get_data(as_text=True)
    assert "PROGRAMA DE OBJETIVOS" in html and "QDV-RG-PG-02_01" in html and "ENE" in html
    assert "Objetivo de Producción" not in html and "Tarea uno" not in html
    # Ninguna celda de seguimiento con datos ni editable (la leyenda de colores sí se muestra).
    assert "data-obj=" not in html and "Editar</button>" not in html and "Realizado" in html
