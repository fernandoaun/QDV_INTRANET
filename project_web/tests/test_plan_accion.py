from __future__ import annotations

import io
import re
from datetime import date, datetime

import pytest
from werkzeug.security import generate_password_hash


def _xlsx() -> bytes:
    """Planilla con la misma estructura que el registro (datos ficticios)."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["C1"], ws["O1"] = "PLAN DE ACCION", "QDV-RG-PG-02_03"
    ws["B7"], ws["D7"] = "RESPONSABLE DE SEGUIMIENTO DE PLAN:", "Responsable test"
    titulos = ["N° ACCIÓN", "ORIGEN  DE LA ACCIÓN", "TIPO DE HALLAZGO RESULTANTE DE LA ACCIÓN", "FECHA DE DETECCIÓN DE LA ACCIÓN",
               "TIPO DE ACCIÓN", "ACCIÓN DEFINIDA", "CAUZA RAÍZ", "DESCRIPCIÓN DE LA CAUSA RAÍZ",
               "REQUISITO/DOCUMENTO ASOCIADO A LA ACCIÓN", "FECHA DE CUMPLIMIENTO", "ESTADO DE LA ACCIÓN",
               "RESPONSABLE DE CUMPLIMIENTO", "MÉTODO DE LA EFICACIA", "RESULTADO DE LA EFICACIA", "COMENTARIOS"]
    for i, t in enumerate(titulos, start=1):
        ws.cell(11, i, t)
    ws.append(["1", "Evaluación ", "Otros", datetime(2026, 9, 30), "Acción de control", "Hacer el procedimiento A.",
               "No aplica", None, "ISO 9001:2015 8.4 ", "31/09/2026", "Atrasado", "Responsable A", None, None, "Atrasado."])
    ws.append(["2", "auditoria interna", "No Conformidad", datetime(2026, 10, 1), "Acción correctiva", "Corregir B.",
               "Aplica", "Falta de control.", "ISO 9001:2015 10.2", datetime(2026, 10, 31), "En curso", "Responsable B"])
    ws.append(["3"])  # fila numerada sin acción: no se carga
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _login(app, client, username, **kw):
    from app.extensions import db
    from app.models import User

    with app.app_context():
        db.session.add(User(username=username, password_hash=generate_password_hash("pw-plan"), activo=True, **kw))
        db.session.commit()
    tok = re.search(r'name="csrf_token"\s+value="([^"]+)"', client.get("/login").get_data(as_text=True)).group(1)
    assert client.post("/login", data={"username": username, "password": "pw-plan", "csrf_token": tok}).status_code in (302, 303)
    return client


def test_leer_planilla(app):
    from app.services import plan_accion_service as svc

    data, avisos = svc.leer_xlsx(_xlsx())
    assert data["responsable_seguimiento"] == "Responsable test"
    a1, a2 = data["acciones"]
    assert a1["numero"] == 1 and a1["accion"] == "Hacer el procedimiento A." and a1["fecha_cumplimiento"] == date(2026, 9, 30)
    assert any("31/09/2026" in x for x in avisos)
    assert a2["numero"] == 2 and a2["causa_raiz_descripcion"] == "Falta de control."


@pytest.fixture
def accion_id(app):
    from app.extensions import db
    from app.services import plan_accion_service as svc

    with app.app_context():
        data, _ = svc.leer_xlsx(_xlsx())
        ids = []
        for fila in data["acciones"]:
            a, errs = svc.crear(fila, None, numero=fila["numero"])
            assert not errs
            ids.append(a.id)
        db.session.commit()
        a = svc.get(ids[0])
        assert a.origen == "Evaluación" and a.fecha_deteccion == date(2026, 9, 30)
        assert svc.get(ids[1]).origen == "Auditoría interna"  # normalizado contra la lista
        return ids[0]


def test_plan_accion_y_en_blanco(app, client, accion_id):
    c = _login(app, client, "pytest_plan_admin", is_admin=True, rol="operaciones")
    html = c.get("/sgi/plan-accion/").get_data(as_text=True)
    for txt in ("PLAN DE ACCIÓN", "QDV-RG-PG-02_03", "IDENTIFICACIÓN Y TRAZABILIDAD", "SEGUIMIENTO", "Hacer el procedimiento A.",
                "30/09/2026", "Atrasado", "Corregir B.", "Nueva acción", "Responsable SGC"):
        assert txt in html, txt
    blanco = c.get("/sgi/plan-accion/en-blanco").get_data(as_text=True)
    assert "QDV-RG-PG-02_03" in blanco and "ACCIÓN DEFINIDA" in blanco
    assert "Hacer el procedimiento A." not in blanco and "Nueva acción" not in blanco


def test_sgi_carga_edita_y_queda_historial(app, client, accion_id):
    from app.services import plan_accion_service as svc

    c = _login(app, client, "pytest_plan_sgi", is_admin=False, rol="sgi")
    r = c.post(f"/sgi/plan-accion/{accion_id}", data={"accion": "Hacer el procedimiento A.", "estado": "cerrado",
                                                       "fecha_cumplimiento": "15/10/2026", "resultado_eficacia": "Eficaz"})
    assert r.status_code in (302, 303)
    with app.app_context():
        a = svc.get(accion_id)
        assert a.estado == "Cerrado" and a.fecha_cumplimiento == date(2026, 10, 15) and a.resultado_eficacia == "Eficaz"
        assert {x.campo for x in svc.cambios(a)} >= {"estado", "fecha_cumplimiento", "resultado_eficacia"}
        assert svc.ultima_actualizacion() is not None
    hist = c.get(f"/sgi/plan-accion/{accion_id}?historial=1").get_data(as_text=True)
    assert "Historial de cambios" in hist and "15/10/2026" in hist
    assert c.post(f"/sgi/plan-accion/{accion_id}", data={"accion": "X", "fecha_deteccion": "99/99/2026"}).status_code in (302, 303)
    with app.app_context():
        assert svc.get(accion_id).accion == "Hacer el procedimiento A."  # fecha inválida: no guarda nada
    r = c.post("/sgi/plan-accion/nueva", data={"accion": "Acción nueva.", "origen": "Incidente"})
    assert r.status_code in (302, 303)
    with app.app_context():
        assert [a.numero for a in svc.listar()] == [1, 2, 3]
        c.post(f"/sgi/plan-accion/{accion_id}/baja")  # la baja es del administrador
        assert svc.get(accion_id) is not None
    c.post("/sgi/plan-accion/responsable", data={"responsable_seguimiento": "Responsable nuevo"})
    assert "Responsable nuevo" in c.get("/sgi/plan-accion/").get_data(as_text=True)


def test_admin_da_de_baja(app, client, accion_id):
    c = _login(app, client, "pytest_plan_admin2", is_admin=True, rol="operaciones")
    c.post(f"/sgi/plan-accion/{accion_id}/baja")
    assert c.get(f"/sgi/plan-accion/{accion_id}").status_code == 404
    assert "Hacer el procedimiento A." not in c.get("/sgi/plan-accion/").get_data(as_text=True)


def test_sin_acceso_sgi(app, client, accion_id):
    c = _login(app, client, "pytest_plan_oper", is_admin=False, rol="operaciones")
    assert c.get("/sgi/plan-accion/", follow_redirects=False).status_code in (302, 303)
    assert c.post(f"/sgi/plan-accion/{accion_id}", data={"accion": "X"}).status_code in (302, 303)


def test_registro_sgi(app):
    from app.services import sgi_procedimiento_service as proc

    with app.app_context():
        links = proc.registro_modulo_links("plan_accion")
    assert links == {"label": "Plan de acción", "blank_url": "/sgi/plan-accion/en-blanco", "filled_url": "/sgi/plan-accion/"}
