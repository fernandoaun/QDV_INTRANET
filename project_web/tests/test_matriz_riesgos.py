from __future__ import annotations

import io
import re
from datetime import datetime

import pytest
from werkzeug.security import generate_password_hash


def _login(app, client, username, **kw):
    from app.extensions import db
    from app.models import User

    with app.app_context():
        db.session.add(User(username=username, password_hash=generate_password_hash("pw-mrg"), activo=True, **kw))
        db.session.commit()
    tok = re.search(r'name="csrf_token"\s+value="([^"]+)"', client.get("/login").get_data(as_text=True)).group(1)
    assert client.post("/login", data={"username": username, "password": "pw-mrg", "csrf_token": tok}).status_code in (302, 303)
    return client


@pytest.fixture
def matriz_id(app):
    from app.extensions import db
    from app.services import matriz_riesgos_service as svc

    with app.app_context():
        m, errs = svc.crear(
            {"sector": "Producción", "filas": [
                {"tipo": "operacional", "riesgo": "Falla de equipos críticos", "probabilidad": "3", "severidad": "5",
                 "acciones": "Cumplir el plan de mantenimiento", "fecha": "12-2026", "responsables": "Responsable de Planta"},
                {"tipo": "Estratégico", "riesgo": "Mejora del rendimiento", "probabilidad": "2", "severidad": "3"},
            ]},
            None,
        )
        assert not errs
        db.session.commit()
        fs = svc.filas(m)
        assert fs[0]["tipo"] == "Operacional" and fs[0]["fecha"] == "12/2026"
        assert svc.resumen(m) == {"Alto": 1, "Moderado": 1}
        return m.id


def test_clasificacion_igual_que_la_planilla():
    from app.services.matriz_riesgos_service import clasificacion

    assert [clasificacion(v) for v in (0, 1, 4, 5, 9, 10, 16, 20, 25)] == [
        "Sin evaluar", "Aceptable", "Bajo", "Moderado", "Moderado", "Alto", "Alto", "Intolerable", "Intolerable"]


def test_listado_hoja_valoracion_y_en_blanco(app, client, matriz_id):
    c = _login(app, client, "pytest_mrg_admin", is_admin=True, rol="operaciones")
    html = c.get("/sgi/matriz-riesgos/").get_data(as_text=True)
    for txt in ("QDV-RG-PG-09_04", "Producción", "Nuevo sector", "mrg-c-alto"):
        assert txt in html, txt
    hoja = c.get(f"/sgi/matriz-riesgos/{matriz_id}").get_data(as_text=True)
    for txt in ("MATRIZ DE RIESGOS Y OPORTUNIDADES", "Identificación de Riesgos", "Falla de equipos críticos", ">15<", "ALTO",
                "MODERADO", "12/2026", "Evaluación de la eficacia de las acciones", "21/07/2026"):
        assert txt in hoja, txt
    val = c.get("/sgi/matriz-riesgos/valoracion").get_data(as_text=True)
    for txt in ("TABLA DE SEVERIDAD", "TABLA DE PROBABILIDAD", "INTOLERABLE", "Riesgo/Impacto de 10 a 16"):
        assert txt in val, txt
    blanco = c.get("/sgi/matriz-riesgos/en-blanco").get_data(as_text=True)
    assert "QDV-RG-PG-09_04" in blanco and "Falla de equipos" not in blanco
    assert "<input" not in blanco.split('class="mrg-sheet"')[1]


def test_sgi_edita_con_historial_y_solo_admin_da_de_baja(app, client, matriz_id):
    from app.services import matriz_riesgos_service as svc

    c = _login(app, client, "pytest_mrg_sgi", is_admin=False, rol="sgi")
    form = {"sector": "Producción", "filas_filas": "2",
            "filas_tipo_0": "Operacional", "filas_riesgo_0": "Falla de equipos críticos", "filas_probabilidad_0": "2",
            "filas_severidad_0": "5", "filas_acciones_0": "Cumplir el plan", "filas_fecha_0": "01/12/2026",
            "filas_responsables_0": "Responsable de Planta", "filas_eficacia_0": "Eficaz"}
    r = c.post(f"/sgi/matriz-riesgos/{matriz_id}", data=form)
    assert r.status_code in (302, 303)
    with app.app_context():
        m = svc.get(matriz_id)
        assert len(svc.filas(m)) == 1 and svc.filas(m)[0]["eficacia"] == "Eficaz"
        assert m.fecha_actualizacion == datetime.now().date()
        h = svc.cambios(m)
        assert h[0].campo == "filas" and "P2×S5 = 10 Alto" in h[0].despues
    form["filas_probabilidad_0"] = "7"
    c.post(f"/sgi/matriz-riesgos/{matriz_id}", data=form)
    with app.app_context():
        assert svc.filas(svc.get(matriz_id))[0]["probabilidad"] == "2"
    c.post(f"/sgi/matriz-riesgos/{matriz_id}/baja", data={})
    with app.app_context():
        assert svc.get(matriz_id) is not None


def test_sector_no_se_repite_y_admin_da_de_baja(app, client, matriz_id):
    from app.extensions import db
    from app.services import matriz_riesgos_service as svc

    with app.app_context():
        m, errs = svc.crear({"sector": "producción"}, None)
        assert m is None and "Ya hay una matriz" in errs[0]
        db.session.rollback()
    c = _login(app, client, "pytest_mrg_admin2", is_admin=True, rol="operaciones")
    c.post(f"/sgi/matriz-riesgos/{matriz_id}/baja", data={})
    with app.app_context():
        assert svc.get(matriz_id) is None


def test_leer_xlsx_una_hoja_por_sector():
    import openpyxl

    from app.services.matriz_riesgos_service import leer_xlsx

    wb = openpyxl.Workbook()
    wb.active.title = "Tabla de valoración"
    wb.active["A1"] = "TABLA DE SEVERIDAD"
    for sector, fecha in (("Gerencia", "12-2026"), ("Producción", datetime(2026, 12, 1))):
        ws = wb.create_sheet(sector)
        ws["K1"] = "QDV-RG-PG-09_04        "
        ws["C6"] = sector
        ws["C7"] = datetime(2026, 9, 29)
        ws.append([])
        for col, val in zip("ABCDEFGHIJK", (1, "Estratégico", f"Riesgo de {sector}", 3, 4, 12, "ALTO", "Acción", fecha, "Resp", None)):
            ws[f"{col}12"] = val
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    hojas = leer_xlsx(buf)
    assert [h["sector"] for h in hojas] == ["Gerencia", "Producción"]
    assert hojas[0]["filas"][0]["fecha"] == "12/2026"
    assert hojas[0]["fecha_actualizacion"].isoformat() == "2026-09-29"
    assert hojas[1]["filas"] == [{"tipo": "Estratégico", "riesgo": "Riesgo de Producción", "probabilidad": "3", "severidad": "4",
                                  "acciones": "Acción", "fecha": "12/2026", "responsables": "Resp", "eficacia": ""}]
